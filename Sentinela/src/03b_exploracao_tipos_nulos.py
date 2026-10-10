"""Fase 1 - tipos e vazios por coluna nas bases que respondem as perguntas (recorte MT).

Uso:  python src/03b_exploracao_tipos_nulos.py
Grava src/outputs/exploracao_tipos_nulos.csv (leva alguns minutos por causa das planilhas da Trajetória).
"""
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
DADOS = RAIZ / "Dados"
SAIDA = RAIZ / "src" / "outputs" / "exploracao_tipos_nulos.csv"
UF_MT = "51"

COLUNAS_CENSO = [
    "NU_ANO_CENSO", "CO_UF", "CO_MUNICIPIO", "CO_IES", "CO_CURSO", "NO_CURSO",
    "CO_CINE_ROTULO", "CO_CINE_AREA_GERAL", "TP_REDE", "TP_CATEGORIA_ADMINISTRATIVA",
    "TP_MODALIDADE_ENSINO", "TP_GRAU_ACADEMICO", "QT_VG_TOTAL", "QT_ING", "QT_MAT", "QT_CONC",
    "QT_SIT_DESVINCULADO", "QT_SIT_TRANCADA", "QT_ING_FIES", "QT_MAT_FIES",
    "QT_MAT_PROUNII", "QT_MAT_PROUNIP",
]
INTEIRO = re.compile(r"-?\d+")
DECIMAL = re.compile(r"-?\d+[.,]\d+")


def tipo_e_estranhos(coluna, valores):
    """Classifica a coluna pelo conteúdo e devolve os valores que fogem do tipo."""
    if valores.empty:
        return "sem valor", ""
    inteiro = valores.str.fullmatch(INTEIRO)
    numero = inteiro | valores.str.fullmatch(DECIMAL)
    if numero.mean() < 0.5:  # maioria numérica: o resto ('SC', '-') é listado como fora do tipo
        return "texto", ""
    estranhos = ", ".join(f"{v!r} ({n})" for v, n in valores[~numero].value_counts().head(3).items())
    if coluna.startswith(("CO_", "Código")):
        return "código (guardar como texto)", estranhos
    if coluna.startswith(("TP_", "IN_")):
        return "categoria codificada", estranhos
    return ("inteiro" if inteiro[numero].all() else "decimal"), estranhos


def perfil(base, arquivo, df):
    linhas = []
    for coluna in df.columns:
        serie = df[coluna].astype("string").str.strip()
        vazios = int((serie.isna() | (serie == "")).sum())
        valores = serie[serie.notna() & (serie != "")]
        tipo, estranhos = tipo_e_estranhos(coluna, valores)
        linhas.append({
            "base": base, "arquivo": arquivo, "coluna": coluna, "linhas_mt": len(df),
            "vazios": vazios, "pct_vazios": round(100 * vazios / len(df), 1) if len(df) else None,
            "distintos": valores.nunique(), "tipo": tipo, "fora_do_tipo": estranhos,
            "com_aspas": int(valores.str.contains('"').sum()),
            "exemplo": valores.iloc[0] if len(valores) else "",
        })
    return linhas


def censo():
    for arquivo in sorted(DADOS.rglob("MICRODADOS_CADASTRO_CURSOS_*.CSV")):
        if arquivo.stem[-4:] < "2021":
            continue
        df = pd.read_csv(arquivo, sep=";", encoding="latin-1", dtype=str,
                         usecols=lambda c: c in COLUNAS_CENSO)
        yield "Censo cursos", arquivo.name, df[df["CO_UF"] == UF_MT]


def trajetoria():
    for arquivo in sorted(DADOS.rglob("indicadores_trajetoria_*.xlsx")):
        wb = load_workbook(arquivo, read_only=True)
        linhas = wb.worksheets[0].iter_rows(min_row=9, values_only=True)
        colunas = [str(c).strip() for c in next(linhas)]
        i_uf = colunas.index("CO_UF")
        mt = [[None if v is None else str(v) for v in l] for l in linhas if str(l[i_uf]) == UF_MT]
        wb.close()
        yield "Trajetória", arquivo.name, pd.DataFrame(mt, columns=colunas)


def qualidade():
    pasta = DADOS / "INEP" / "indicadores_qualidade"
    arquivos = sorted(pasta.glob("[Cc][Pp][Cc]_*.xlsx")) + sorted(pasta.glob("[Ii][Gg][Cc]_*.xlsx"))
    arquivos.append(pasta / "conceito_enade_licenciaturas.xlsx")
    for arquivo in arquivos:
        df = pd.read_excel(arquivo, sheet_name=0, dtype=str)
        df.columns = [re.sub(r"[*¹²]", "", str(c)).strip() for c in df.columns]
        col_uf = next(c for c in df.columns if c.startswith("Sigla da UF"))
        base = "IGC" if "igc" in arquivo.name.lower() else "CPC" if "cpc" in arquivo.name.lower() else "Enade Licenciaturas"
        yield base, arquivo.name, df[df[col_uf].str.strip() == "MT"]


def main():
    linhas = []
    for leitor in (censo, qualidade, trajetoria):
        for base, arquivo, df in leitor():
            linhas += perfil(base, arquivo, df)
            print(f"lido: {arquivo} ({len(df)} linhas de MT)", flush=True)
    resultado = pd.DataFrame(linhas)
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    resultado.to_csv(SAIDA, index=False)

    print("\n== colunas com vazio, valor fora do tipo ou aspas ==")
    alerta = resultado[(resultado["vazios"] > 0) | (resultado["fora_do_tipo"] != "") | (resultado["com_aspas"] > 0)]
    print(alerta[["arquivo", "coluna", "linhas_mt", "vazios", "pct_vazios", "tipo", "fora_do_tipo", "com_aspas"]].to_string(index=False))
    print("\n== tipo por coluna (primeiro arquivo de cada base) ==")
    primeiro = resultado.drop_duplicates(["base", "coluna"])
    print(primeiro[["base", "coluna", "tipo", "distintos", "exemplo"]].to_string(index=False))


if __name__ == "__main__":
    main()
