"""Fase 2 - carga das bases no PostgreSQL (schema raw): as de P1 a P5, as de apoio e o IBGE do B1.

Uso:  python src/04_carga_raw.py                 carrega tudo (cópia fiel, Brasil inteiro)
      python src/04_carga_raw.py --so-mt         carrega só as linhas de Mato Grosso (mais rápido)
      python src/04_carga_raw.py --fontes cpc enade_licenciaturas
      python src/04_carga_raw.py --simular       lê os arquivos e conta, sem tocar no banco

Cada arquivo entra em uma transação que apaga as linhas daquele arquivo e insere de novo:
rodar duas vezes dá a mesma contagem. A contagem fica em src/outputs/carga_raw_contagem.csv.
O acesso ao banco vem do .env (PGHOST, PGPORT, PGDATABASE, PGUSER, PGPASSWORD).
"""
import argparse
import csv
import json
import re
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
DADOS = RAIZ / "Dados"
CONTAGEM = RAIZ / "src" / "outputs" / "carga_raw_contagem.csv"
ANOS_CENSO = ("2021", "2022", "2023", "2024")  # a P3 pede 2021 a 2024

# tabela em raw -> onde estão os arquivos, como ler e que pergunta usa
FONTES = {
    "trajetoria": {
        "perguntas": "P1, P2, P4, P5",
        "pasta": "INEP/trajetoria",
        "padrao": "indicadores_trajetoria_*.xlsx",
        "linha_cabecalho": 9,  # os nomes técnicos (CO_CURSO, TDA...) estão na linha 9
        "filtro_mt": ("co_uf", "51"),
    },
    "censo_cursos": {
        "perguntas": "P3",
        "pasta": "INEP/censo_superior",
        "padrao": "MICRODADOS_CADASTRO_CURSOS_*.CSV",
        "anos": ANOS_CENSO,
        "filtro_mt": ("co_uf", "51"),
    },
    "cpc": {
        "perguntas": "P4",
        "pasta": "INEP/indicadores_qualidade",
        "padrao": "[Cc][Pp][Cc]_*.xlsx",
        "linha_cabecalho": 1,
        "filtro_mt": ("sigla_da_uf", "MT"),
    },
    "enade_licenciaturas": {
        "perguntas": "P5",
        "pasta": "INEP/indicadores_qualidade",
        "padrao": "conceito_enade_licenciaturas.xlsx",
        "linha_cabecalho": 1,
        "filtro_mt": ("sigla_da_uf", "MT"),
    },
    "censo_ies": {
        "perguntas": "apoio (sede da IES)",
        "pasta": "INEP/censo_superior",
        "padrao": "MICRODADOS_*IES_*.CSV",  # o nome muda de CADASTRO_IES para ED_SUP_IES em 2022
        "anos": ANOS_CENSO,
        "filtro_mt": ("co_uf_ies", "51"),
    },
    "igc": {
        "perguntas": "apoio (qualidade da IES)",
        "pasta": "INEP/indicadores_qualidade",
        "padrao": "[Ii][Gg][Cc]_*.xlsx",
        "linha_cabecalho": 1,
        "filtro_mt": ("sigla_da_uf", "MT"),
    },
    "idd": {
        "perguntas": "apoio (P4)",
        "pasta": "INEP/indicadores_qualidade",
        "padrao": "[Ii][Dd][Dd]_*.xlsx",
        "linha_cabecalho": 1,
        "filtro_mt": ("sigla_da_uf", "MT"),
    },
    "enade": {
        "perguntas": "apoio (P4)",
        "pasta": "INEP/indicadores_qualidade",
        "padrao": "conceito_enade_20*.xlsx",
        "linha_cabecalho": 1,
        "filtro_mt": ("sigla_da_uf", "MT"),
    },
    "ibge_populacao": {
        "perguntas": "B1",
        "pasta": "IBGE/sidra",
        "padrao": "tabela_9514_populacao_*.json",
        "filtro_mt": ("co_uf", "51"),
    },
}


def nome_coluna(texto):
    """' Código do Curso**' -> 'codigo_do_curso': sem acento, sem asterisco, minúsculo."""
    sem_acento = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", sem_acento.lower()).strip("_")[:60] or "coluna"


def nomes_colunas(cabecalho):
    """Normaliza o cabeçalho e numera os nomes que ficarem repetidos."""
    nomes, vistos = [], {}
    for texto in cabecalho:
        nome = nome_coluna(texto)
        vistos[nome] = vistos.get(nome, 0) + 1
        nomes.append(nome if vistos[nome] == 1 else f"{nome}_{vistos[nome]}")
    return nomes


def ler_csv(arquivo):
    """CSV do Censo: separador ; e latin-1. Devolve (colunas, gerador de linhas)."""
    f = open(arquivo, encoding="latin-1", newline="")
    leitor = csv.reader(f, delimiter=";")
    colunas = nomes_colunas(next(leitor))

    def linhas():
        with f:
            for linha in leitor:
                yield [v if v != "" else None for v in linha]

    return colunas, linhas()


def ler_xlsx(arquivo, linha_cabecalho):
    """Planilha do INEP, lida em fluxo. Usa a primeira aba que tem cabeçalho de verdade."""
    wb = load_workbook(arquivo, read_only=True)
    for ws in wb.worksheets:
        it = ws.iter_rows(min_row=linha_cabecalho, values_only=True)
        cabecalho = [c for c in next(it, ()) if c is not None]
        if len(cabecalho) >= 5:  # algumas planilhas trazem uma aba de aviso antes dos dados
            break
    else:
        raise ValueError(f"{arquivo.name}: nenhuma aba com cabeçalho na linha {linha_cabecalho}")
    colunas = nomes_colunas(cabecalho)

    def linhas():
        for linha in it:
            if linha[0] is None:  # linhas em branco e notas de rodapé
                continue
            yield [None if v is None else str(v) for v in linha[:len(colunas)]]
        wb.close()

    return colunas, linhas()


def ler_json_sidra(arquivo):
    """JSON da API de agregados do IBGE: achata variável -> classificações -> séries em linhas."""
    dados = json.loads(arquivo.read_text(encoding="utf-8"))
    classificacoes = [c["nome"] for c in dados[0]["resultados"][0]["classificacoes"]]
    colunas = nomes_colunas(["variavel_id", "variavel", "unidade", *classificacoes,
                             "co_uf", "co_municipio", "localidade", "nivel", "ano", "valor"])

    def linhas():
        for variavel in dados:
            for resultado in variavel["resultados"]:
                categorias = [next(iter(c["categoria"].values())) for c in resultado["classificacoes"]]
                for serie in resultado["series"]:
                    local = serie["localidade"]
                    for ano, valor in serie["serie"].items():
                        yield [variavel["id"], variavel["variavel"], variavel["unidade"], *categorias,
                               local["id"][:2], local["id"], local["nome"], local["nivel"]["id"], ano, valor]

    return colunas, linhas()


def arquivos_da_fonte(fonte):
    arquivos = sorted((DADOS / fonte["pasta"]).glob(fonte["padrao"]))
    if "anos" in fonte:
        arquivos = [a for a in arquivos if a.stem[-4:] in fonte["anos"]]
    return arquivos


def ler(arquivo, fonte):
    if arquivo.suffix.lower() == ".csv":
        return ler_csv(arquivo)
    if arquivo.suffix.lower() == ".json":
        return ler_json_sidra(arquivo)
    return ler_xlsx(arquivo, fonte["linha_cabecalho"])


def filtrar(colunas, linhas, fonte, so_mt, contador):
    """Conta as linhas lidas e, com --so-mt, deixa passar só as de Mato Grosso."""
    coluna, valor = fonte["filtro_mt"]
    i = colunas.index(coluna)
    n = len(colunas)
    for linha in linhas:
        contador["lidas"] += 1
        if so_mt and (linha[i] or "").strip() != valor:
            continue
        contador["carregadas"] += 1
        yield (linha + [None] * n)[:n]


def preparar_tabela(conn, tabela, colunas):
    """Cria a tabela se não existir e acrescenta as colunas que este arquivo trouxer a mais."""
    from psycopg import sql

    alvo = sql.Identifier("raw", tabela)
    conn.execute("CREATE SCHEMA IF NOT EXISTS raw")
    conn.execute(sql.SQL(
        "CREATE TABLE IF NOT EXISTS {} ("
        "arquivo_origem text NOT NULL, carregado_em timestamptz NOT NULL DEFAULT now())"
    ).format(alvo))
    for coluna in colunas:
        conn.execute(sql.SQL("ALTER TABLE {} ADD COLUMN IF NOT EXISTS {} text").format(
            alvo, sql.Identifier(coluna)))
    conn.execute(sql.SQL("CREATE INDEX IF NOT EXISTS {} ON {} (arquivo_origem)").format(
        sql.Identifier(f"ix_{tabela}_arquivo"), alvo))


def carregar_arquivo(conn, tabela, arquivo, colunas, linhas):
    """Apaga o que já veio deste arquivo e insere de novo, tudo na mesma transação."""
    from psycopg import sql

    alvo = sql.Identifier("raw", tabela)
    conta = sql.SQL("SELECT count(*) FROM {} WHERE arquivo_origem = %s").format(alvo)
    copia = sql.SQL("COPY {} (arquivo_origem, {}) FROM STDIN").format(
        alvo, sql.SQL(", ").join(sql.Identifier(c) for c in colunas))
    with conn.transaction():
        preparar_tabela(conn, tabela, colunas)
        antes = conn.execute(conta, [arquivo.name]).fetchone()[0]
        conn.execute(sql.SQL("DELETE FROM {} WHERE arquivo_origem = %s").format(alvo), [arquivo.name])
        with conn.cursor().copy(copia) as copy:
            for linha in linhas:
                copy.write_row([arquivo.name] + linha)
        depois = conn.execute(conta, [arquivo.name]).fetchone()[0]
    return antes, depois


def conectar():
    import psycopg
    from dotenv import load_dotenv

    load_dotenv(RAIZ / ".env")  # o psycopg lê PGHOST, PGUSER... do ambiente
    return psycopg.connect(autocommit=True)


def gravar_contagem(registros):
    CONTAGEM.parent.mkdir(parents=True, exist_ok=True)
    novo = not CONTAGEM.exists()
    with open(CONTAGEM, "a", encoding="utf-8", newline="") as f:
        escritor = csv.DictWriter(f, fieldnames=list(registros[0]))
        if novo:
            escritor.writeheader()
        escritor.writerows(registros)


def main():
    p = argparse.ArgumentParser(description="Carga das bases no schema raw.")
    p.add_argument("--fontes", nargs="+", choices=list(FONTES), default=list(FONTES))
    p.add_argument("--so-mt", action="store_true", help="carrega só as linhas de Mato Grosso")
    p.add_argument("--simular", action="store_true", help="lê e conta, sem tocar no banco")
    args = p.parse_args()

    conn = None if args.simular else conectar()
    executado_em = datetime.now().isoformat(timespec="seconds")
    registros = []
    for tabela in args.fontes:
        fonte = FONTES[tabela]
        arquivos = arquivos_da_fonte(fonte)
        if not arquivos:
            print(f"raw.{tabela}: nenhum arquivo em Dados/{fonte['pasta']} ({fonte['padrao']})")
            continue
        for arquivo in arquivos:
            colunas, linhas = ler(arquivo, fonte)
            contador = {"lidas": 0, "carregadas": 0}
            linhas = filtrar(colunas, linhas, fonte, args.so_mt, contador)
            if args.simular:
                for _ in linhas:
                    pass
                antes = depois = None
            else:
                antes, depois = carregar_arquivo(conn, tabela, arquivo, colunas, linhas)
            registros.append({
                "executado_em": executado_em, "tabela": f"raw.{tabela}", "perguntas": fonte["perguntas"],
                "arquivo": arquivo.name, "colunas": len(colunas), "linhas_lidas": contador["lidas"],
                "linhas_carregadas": contador["carregadas"], "so_mt": args.so_mt,
                "no_banco_antes": antes, "no_banco_depois": depois,
            })
            print(f"raw.{tabela:<20} {arquivo.name:<58} {len(colunas):>3} colunas  "
                  f"{contador['lidas']:>7} lidas  {contador['carregadas']:>7} carregadas"
                  + ("" if args.simular else f"  (no banco: {antes} -> {depois})"), flush=True)

    if conn is not None:
        conn.close()
        if registros:
            gravar_contagem(registros)
            print(f"contagem gravada em {CONTAGEM.relative_to(RAIZ)}")


if __name__ == "__main__":
    main()
