"""Fase 1 - estrutura das planilhas do INEP (Trajetória, CPC, Enade, IDD, IGC) e do JSON do IBGE.

Uso:  python src/02_exploracao_planilhas.py
Mostra, para cada arquivo, as abas, em que linha está o cabeçalho e os nomes das colunas.
"""
import json
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
DADOS = RAIZ / "Dados"


def estrutura_xlsx(arquivo, linhas_topo=12):
    wb = load_workbook(arquivo, read_only=True)
    print(f"\n=== {arquivo.relative_to(DADOS)} ({arquivo.stat().st_size / 1e6:.1f} MB)")
    for ws in wb.worksheets:
        print(f"  aba '{ws.title}': {ws.max_row} linhas x {ws.max_column} colunas")
        for n, linha in enumerate(ws.iter_rows(min_row=1, max_row=linhas_topo, values_only=True), start=1):
            preenchidas = [str(v).replace("\n", " ")[:28] for v in linha if v is not None]
            # o cabeçalho é a primeira linha com muitas células preenchidas
            if len(preenchidas) >= 5:
                print(f"    linha {n} ({len(preenchidas)} células): {preenchidas}")
    wb.close()


def estrutura_json(arquivo):
    dados = json.loads(arquivo.read_text(encoding="utf-8"))
    print(f"\n=== {arquivo.relative_to(DADOS)} ({arquivo.stat().st_size / 1e6:.1f} MB)")
    print(f"  tipo: {type(dados).__name__}, itens: {len(dados)}")
    print(f"  amostra: {json.dumps(dados[:2] if isinstance(dados, list) else dados, ensure_ascii=False)[:900]}")


def main():
    for arquivo in sorted(DADOS.rglob("*.xlsx")):
        estrutura_xlsx(arquivo)
    for arquivo in sorted(DADOS.rglob("*.json")):
        estrutura_json(arquivo)


if __name__ == "__main__":
    main()
