"""Fase 1 - perfil dos Indicadores de Trajetória (recorte MT) e taxa de junção com Censo, CPC e Enade.

Uso:  python src/03_exploracao_trajetoria.py
Grava src/outputs/exploracao_trajetoria.csv e src/outputs/exploracao_juncoes.csv.
"""
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
DADOS = RAIZ / "Dados"
SAIDAS = RAIZ / "src" / "outputs"
UF_MT = "51"
LINHA_CABECALHO = 9  # os nomes técnicos (CO_CURSO, TDA...) estão na linha 9


def ler_trajetoria_mt(arquivo):
    """Lê a planilha em fluxo e devolve (total de linhas, DataFrame só de MT, tudo como texto)."""
    wb = load_workbook(arquivo, read_only=True)
    linhas = wb.worksheets[0].iter_rows(min_row=LINHA_CABECALHO, values_only=True)
    colunas = [str(c).strip() for c in next(linhas)]
    i_uf = colunas.index("CO_UF")
    total, mt = 0, []
    for linha in linhas:
        if linha[0] is None:
            continue
        total += 1
        if str(linha[i_uf]) == UF_MT:
            mt.append([None if v is None else str(v) for v in linha])
    wb.close()
    return total, pd.DataFrame(mt, columns=colunas)


def perfil(arquivo, total, df):
    num = df[["QT_INGRESSANTE", "TAP", "TCA", "TDA"]].apply(pd.to_numeric, errors="coerce")
    chave = ["CO_CURSO", "NU_ANO_INGRESSO", "NU_ANO_REFERENCIA"]
    por_curso = num["QT_INGRESSANTE"].groupby(df["CO_CURSO"]).first()
    soma = num["TAP"] + num["TCA"] + num["TDA"]
    return {
        "arquivo": arquivo.name,
        "linhas": total,
        "linhas_mt": len(df),
        "cursos_mt": df["CO_CURSO"].nunique(),
        "ies_mt": df["CO_IES"].nunique(),
        "municipios_mt": df["CO_MUNICIPIO"].nunique(),
        "ano_ingresso": "/".join(sorted(df["NU_ANO_INGRESSO"].unique())),
        "ano_referencia": f"{df['NU_ANO_REFERENCIA'].min()}-{df['NU_ANO_REFERENCIA'].max()}",
        "chave_repetida": int(df.duplicated(chave).sum()),
        "cursos_presencial": df.loc[df["TP_MODALIDADE_ENSINO"] == "1", "CO_CURSO"].nunique(),
        "cursos_ead": df.loc[df["TP_MODALIDADE_ENSINO"] == "2", "CO_CURSO"].nunique(),
        "ingressantes_mt": int(por_curso.sum()),
        "cursos_menos_de_10_ingressantes": int((por_curso < 10).sum()),
        "linhas_tap_tca_tda_fora_de_100": int(((soma - 100).abs() > 0.5).sum()),
        "nulos_em_tda": int(num["TDA"].isna().sum()),
    }


def cursos_censo_mt():
    cursos = {}
    for arquivo in sorted(DADOS.rglob("MICRODADOS_CADASTRO_CURSOS_*.CSV")):
        ano = arquivo.stem[-4:]
        df = pd.read_csv(arquivo, sep=";", encoding="latin-1", dtype=str, usecols=["CO_UF", "CO_CURSO"])
        cursos[ano] = set(df.loc[df["CO_UF"] == UF_MT, "CO_CURSO"])
    return cursos


def cursos_planilha_mt(arquivo):
    """Códigos de curso de MT numa planilha de indicador de qualidade (CPC, Enade)."""
    df = pd.read_excel(arquivo, sheet_name=0, dtype=str)
    if len(df.columns) < 5:  # 2021 traz uma aba de aviso antes da aba de dados
        df = pd.read_excel(arquivo, sheet_name=1, dtype=str)
    df.columns = [c.strip().replace("*", "").strip() for c in df.columns]
    col_curso = next(c for c in df.columns if c.startswith("Código do Curso"))
    col_uf = next(c for c in df.columns if c.startswith("Sigla da UF"))
    mt = df[df[col_uf].str.strip() == "MT"]
    return len(df), set(mt[col_curso].str.strip())


def taxa(nome, base, outra):
    achou = len(base & outra)
    return {"juncao": nome, "cursos_na_base": len(base), "cursos_na_outra": len(outra),
            "com_par": achou, "taxa_pct": round(100 * achou / len(base), 1) if base else None}


def main():
    SAIDAS.mkdir(parents=True, exist_ok=True)
    perfis, traj, licenciaturas = [], set(), set()
    for arquivo in sorted(DADOS.rglob("indicadores_trajetoria_*.xlsx")):
        total, df = ler_trajetoria_mt(arquivo)
        perfis.append(perfil(arquivo, total, df))
        traj |= set(df["CO_CURSO"])
        licenciaturas |= set(df.loc[df["TP_GRAU_ACADEMICO"] == "2", "CO_CURSO"])
        print(f"lido: {arquivo.name}", flush=True)
    pd.DataFrame(perfis).to_csv(SAIDAS / "exploracao_trajetoria.csv", index=False)
    print(pd.DataFrame(perfis).T.to_string())

    juncoes = []
    censo = cursos_censo_mt()
    for ano, cursos in censo.items():
        juncoes.append(taxa(f"Trajetória MT x Censo cursos {ano}", traj, cursos))
    recentes = set().union(*[c for a, c in censo.items() if a >= "2021"])
    juncoes.append(taxa("Trajetória MT x Censo cursos 2021-2024 (qualquer ano)", traj, recentes))

    qualidade = DADOS / "INEP" / "indicadores_qualidade"
    cpc = set()
    for arquivo in sorted(qualidade.glob("[Cc][Pp][Cc]_*.xlsx")):
        _, cursos = cursos_planilha_mt(arquivo)
        cpc |= cursos
        juncoes.append(taxa(f"Trajetória MT x {arquivo.name}", traj, cursos))
    juncoes.append(taxa("Trajetória MT x CPC 2021-2023 (qualquer ano)", traj, cpc))
    _, enade_lic = cursos_planilha_mt(qualidade / "conceito_enade_licenciaturas.xlsx")
    juncoes.append(taxa("Licenciaturas da Trajetória MT x Enade Licenciaturas", licenciaturas, enade_lic))
    juncoes.append(taxa("Enade Licenciaturas MT x Trajetória MT", enade_lic, traj))

    pd.DataFrame(juncoes).to_csv(SAIDAS / "exploracao_juncoes.csv", index=False)
    print(pd.DataFrame(juncoes).to_string(index=False))


if __name__ == "__main__":
    main()
